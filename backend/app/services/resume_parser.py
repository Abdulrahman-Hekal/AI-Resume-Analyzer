import io
from fastapi import HTTPException, status
import pypdf
import docx
from app.utils.text_processing import clean_text
from app.core.logging import logger

class ResumeParser:
    """Handles text extraction from PDF and DOCX documents with robust error handling."""

    @classmethod
    def extract_text_from_pdf(cls, file_bytes: bytes) -> str:
        """Extract text from PDF file bytes, handling scanned/empty PDF edge cases."""
        try:
            reader = pypdf.PdfReader(io.BytesIO(file_bytes))
            if len(reader.pages) == 0:
                raise HTTPException(
                    status_code=status.HTTP_400_BAD_REQUEST,
                    detail="The uploaded PDF contains no pages."
                )

            extracted_parts = []
            for idx, page in enumerate(reader.pages):
                page_text = page.extract_text() or ""
                extracted_parts.append(page_text.strip())

            full_text = "\n\n".join(extracted_parts).strip()

            # Check for scanned or image-only PDF
            if len(full_text) < 30:
                raise HTTPException(
                    status_code=status.HTTP_400_BAD_REQUEST,
                    detail="The uploaded PDF appears to be a scanned image or contains no readable text. Please upload a PDF generated with selectable digital text."
                )

            return clean_text(full_text)

        except HTTPException:
            raise
        except pypdf.errors.PdfReadError as e:
            logger.error(f"Corrupted PDF: {e}")
            raise HTTPException(
                status_code=status.HTTP_400_BAD_REQUEST,
                detail="The PDF file is corrupted or protected with an unsupported password."
            )
        except Exception as e:
            logger.error(f"Unexpected error parsing PDF: {e}")
            raise HTTPException(
                status_code=status.HTTP_400_BAD_REQUEST,
                detail=f"Failed to read the PDF document: {str(e)}"
            )

    @classmethod
    def extract_text_from_docx(cls, file_bytes: bytes) -> str:
        """Extract text from DOCX file bytes including paragraphs and table contents."""
        try:
            doc = docx.Document(io.BytesIO(file_bytes))
            text_parts = []

            # Extract from paragraphs
            for p in doc.paragraphs:
                p_text = p.text.strip()
                if p_text:
                    text_parts.append(p_text)

            # Extract from tables (many resumes format skills/education in tables)
            for table in doc.tables:
                for row in table.rows:
                    row_texts = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                    if row_texts:
                        text_parts.append(" | ".join(row_texts))

            full_text = "\n".join(text_parts).strip()

            if len(full_text) < 30:
                raise HTTPException(
                    status_code=status.HTTP_400_BAD_REQUEST,
                    detail="The DOCX document contains insufficient readable text. Please ensure it is a valid, non-empty resume."
                )

            return clean_text(full_text)

        except HTTPException:
            raise
        except Exception as e:
            logger.error(f"Failed to parse DOCX: {e}")
            raise HTTPException(
                status_code=status.HTTP_400_BAD_REQUEST,
                detail="The DOCX file is malformed or corrupted."
            )

    @classmethod
    def parse(cls, file_bytes: bytes, file_type: str) -> str:
        """Dispatcher to parse resume bytes based on file_type ('pdf' or 'docx')."""
        if file_type.lower() == "pdf":
            return cls.extract_text_from_pdf(file_bytes)
        elif file_type.lower() == "docx":
            return cls.extract_text_from_docx(file_bytes)
        else:
            raise HTTPException(
                status_code=status.HTTP_400_BAD_REQUEST,
                detail=f"Unsupported file type: {file_type}"
            )
